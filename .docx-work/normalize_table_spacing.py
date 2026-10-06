from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

SOURCE = Path(r"C:\Users\Tom Pc\Desktop\Commissions\MiniCapstone1\CERTIGUARD_CHAPTER_3_PLAIN_TABLES.docx")
OUTPUT = Path(r"C:\Users\Tom Pc\Desktop\Commissions\MiniCapstone1\CERTIGUARD_CHAPTER_3_FINAL.docx")

def set_cell_margins(cell, top=55, start=65, bottom=55, end=65):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in('w:tcMar')
    if tc_mar is None:
        tc_mar = OxmlElement('w:tcMar')
        tc_pr.append(tc_mar)
    for side, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        element = tc_mar.find(qn(f'w:{side}'))
        if element is None:
            element = OxmlElement(f'w:{side}')
            tc_mar.append(element)
        element.set(qn('w:w'), str(value))
        element.set(qn('w:type'), 'dxa')

doc = Document(SOURCE)

# Remove oversized blank-paragraph gaps generated around tables.
for paragraph in doc.paragraphs:
    if not paragraph.text.strip():
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

for table in doc.tables:
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
                for run in paragraph.runs:
                    run.font.size = Pt(10 if row_index else 10.5)
                    run.bold = row_index == 0

doc.save(OUTPUT)
print(f'Normalized spacing in {len(doc.tables)} tables: {OUTPUT.name}')
