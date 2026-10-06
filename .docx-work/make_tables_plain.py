from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

SOURCE = Path(r"C:\Users\Tom Pc\Desktop\Commissions\MiniCapstone1\CERTIGUARD_CHAPTER_3.docx")
OUTPUT = Path(r"C:\Users\Tom Pc\Desktop\Commissions\MiniCapstone1\CERTIGUARD_CHAPTER_3_PLAIN_TABLES.docx")

def remove_shading(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    for shd in list(tc_pr.findall(qn('w:shd'))):
        tc_pr.remove(shd)

def set_black_borders(table):
    tbl_pr = table._tbl.tblPr
    for old in list(tbl_pr.findall(qn('w:tblBorders'))):
        tbl_pr.remove(old)
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        border = OxmlElement(f'w:{edge}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '8')
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), '000000')
        borders.append(border)
    tbl_pr.append(borders)

def clear_color(run):
    rpr = run._r.get_or_add_rPr()
    for color in list(rpr.findall(qn('w:color'))):
        rpr.remove(color)

doc = Document(SOURCE)
for table in doc.tables:
    set_black_borders(table)
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            remove_shading(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = 'Times New Roman'
                    run._element.rPr.rFonts.set(qn('w:ascii'), 'Times New Roman')
                    run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Times New Roman')
                    run.font.size = Pt(10)
                    run.bold = row_index == 0
                    clear_color(run)

doc.save(OUTPUT)
print(f'Updated {len(doc.tables)} tables.')
