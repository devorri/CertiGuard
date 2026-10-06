from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

SOURCE = Path(r"C:\Users\Tom Pc\Downloads\MINI-CAPSTONE12.docx")
OUTPUT = Path(r"C:\Users\Tom Pc\Desktop\Commissions\MiniCapstone1\CERTIGUARD_CHAPTER_3.docx")

doc = Document(SOURCE)

def source_para(text):
    return next(p for p in doc.paragraphs if p.text.strip() == text)

chapter_style = source_para('CHAPTER I')
body_heading_style = source_para('Background of the Study')
body_style = next(p for p in doc.paragraphs if p.text.startswith('In the Philippines, barangay governance'))

def clone_format(target, template, centered=False, bold=None):
    if template._p.pPr is not None:
        target._p.get_or_add_pPr().append(deepcopy(template._p.pPr))
    if template.runs:
        tr = template.runs[0]
        run = target.add_run()
        if tr._r.rPr is not None:
            run._r.get_or_add_rPr().append(deepcopy(tr._r.rPr))
        if bold is not None:
            run.bold = bold
        return run
    return target.add_run()

def add_paragraph(text='', kind='body'):
    p = doc.add_paragraph()
    if kind == 'chapter':
        r = clone_format(p, chapter_style, bold=True)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif kind == 'heading':
        r = clone_format(p, body_heading_style, bold=True)
    elif kind == 'caption':
        r = clone_format(p, body_style)
        r.italic = True
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        r = clone_format(p, body_style)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r.text = text
    return p

def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tc_pr.append(shd)

def set_cell_margins(cell, top=100, start=100, bottom=100, end=100):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn(f'w:{m}'))
        if node is None:
            node = OxmlElement(f'w:{m}')
            tcMar.append(node)
        node.set(qn('w:w'), str(v)); node.set(qn('w:type'), 'dxa')

def add_table(headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        border = OxmlElement(f'w:{edge}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '4')
        border.set(qn('w:color'), 'D9D9D9')
        borders.append(border)
    tbl_pr.append(borders)
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header
        set_cell_shading(cell, '1F4E78')
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.font.name = 'Arial'; run.font.size = Pt(8); run.font.color.rgb = None; run.bold = True
            rpr = run._r.get_or_add_rPr(); color = OxmlElement('w:color'); color.set(qn('w:val'), 'FFFFFF'); rpr.append(color)
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
            set_cell_margins(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2:
                set_cell_shading(cells[i], 'EAF2F8')
            for run in cells[i].paragraphs[0].runs:
                run.font.name = 'Arial'; run.font.size = Pt(8)
        if widths:
            for i, width in enumerate(widths):
                cells[i].width = Inches(width)
    doc.add_paragraph()
    return table

doc.add_page_break()
add_paragraph('CHAPTER III', 'chapter')
add_paragraph('METHODOLOGY', 'chapter')

add_paragraph('Technicality of the Project', 'heading')
add_paragraph('The proposed system is titled CertiGuard Secure Document Issuance and Verification Framework for Barangay Taguranao. It is a web-based application designed to digitize the registration, certificate-request, approval, issuance, and verification processes of the barangay. The system is intended to improve service delivery while protecting resident information and reducing the opportunity for document forgery.')
add_paragraph('Residents create an account through a browser-based registration form. The form records their basic profile and requires the upload of a test-only dummy valid ID image. The uploaded image is stored separately from certificate data and is available only to authorized barangay staff for residency checking. In the production design, the file is kept in a protected Supabase Storage bucket; no valid-ID image, file path, or image content is included in a generated certificate or public verification result.')
add_paragraph('Before a resident can use the certificate portal, the Barangay Secretary reviews the submitted dummy ID and either approves or rejects the registration. Approved residents may submit a request for Barangay Clearance, Certificate of Indigency, or Certificate of Residency. The request is then placed in the administrative queue for evaluation and processing.')
add_paragraph('The administrative console enables authorized staff to review resident profiles, ID-verification status, certificate applications, and transaction records. Upon approval of a certificate request, the system generates a unique control number and a SHA-256 hash signature. A QR code containing the verification reference is embedded in the downloadable certificate, allowing an employer, school, or agency to confirm the document through the public verification page.')
add_paragraph('The prototype front end is developed with React, TypeScript, HTML, and CSS through Vite. The planned production backend uses Supabase Authentication, PostgreSQL, and private Supabase Storage for protected identity files. The prototype uses simulated accounts and dummy ID images only, following the Data Privacy Act of 2012 and avoiding the collection or display of real IDs during testing.')

add_paragraph('Locale of the Study', 'heading')
add_paragraph('The study is conducted for Barangay Taguranao. The proposed system supports the barangay office, its residents, and authorized third-party institutions that need to verify a certificate. Residents and staff access the system through a web browser using a desktop computer, laptop, tablet, or mobile device with an internet connection. The system is designed to support the existing certificate-issuance functions of the barangay without changing the authority of barangay officials to evaluate and approve official records.')
add_paragraph('Figure 1: Locale of the Study', 'caption')

add_paragraph('Requirements for Modeling', 'heading')
add_paragraph('The system models the information and activities required to register a resident, verify residency, submit a certificate request, issue an authenticated certificate, and validate a certificate presented to an external institution. The primary models used are the Input Process Output Chart, Entity Relationship Diagram, Data Dictionary, Context Diagram, Data Flow Diagram, System Flowchart, and Disaster Recovery Plan.')

add_paragraph('IPO Chart', 'heading')
add_paragraph('The Input Process Output Chart presents how resident and administrative data are transformed into verified certificates and monitoring information. Inputs are submitted by residents and staff, the process applies validation and approval rules, and the output provides certificates, verification results, notifications, and reports.')
add_table(['Input', 'Process', 'Output'], [[
    'Resident profile, dummy valid ID image, certificate type, purpose, staff review decision, certificate records',
    'Register account, validate fields, store ID in protected storage, Secretary review, process request, generate SHA-256 hash and QR code, send notification, verify certificate',
    'Approved resident account, ID-review status, certificate request status, authenticated PDF certificate, QR verification result, SMS notification, administrative reports'
]], [2.15, 2.75, 2.15])
add_paragraph('Figure 2: Input Process Output Chart', 'caption')

add_paragraph('Entity Relationship Diagram', 'heading')
add_paragraph('The Entity Relationship Diagram defines the relationships among the main records used by CertiGuard. A resident account has one ID-verification record and can submit many certificate requests. Each approved certificate request produces one issued certificate. Staff members review the ID-verification record and process certificate requests. Every issued certificate contains a control number and hash signature that are used by the public verification function.')
add_table(['Entity', 'Relationship'], [
    ['Resident', 'uploads one ID Verification record and submits many Certificate Requests'],
    ['ID Verification', 'belongs to one Resident and is reviewed by one Staff member'],
    ['Certificate Request', 'belongs to one Resident and may produce one Issued Certificate'],
    ['Issued Certificate', 'belongs to one Certificate Request and is verified through its control number and hash signature'],
    ['Staff', 'reviews registration IDs and processes Certificate Requests'],
], [2.1, 4.95])
add_paragraph('Figure 3: Entity Relationship Diagram', 'caption')

add_paragraph('Data Dictionary', 'heading')
add_paragraph('The data dictionary identifies the main data records used in the system and establishes the purpose of each field. The ID image is represented only by a protected storage path and metadata. Its content is not copied into certificate or verification records.')
add_table(['Table', 'Key Fields', 'Description'], [
    ['Users', 'user_id, full_name, email, phone, address, purok, role', 'Stores resident and authorized staff profiles. Email is unique; role controls access.'],
    ['ID Verifications', 'verification_id, user_id, storage_path, file_name, status, reviewed_by, reviewed_at', 'Stores private dummy-ID upload metadata and the Secretary review decision.'],
    ['Certificate Requests', 'request_id, user_id, control_number, type, purpose, status, processed_by', 'Stores resident applications for barangay certificates and their processing state.'],
    ['Issued Certificates', 'certificate_id, request_id, control_number, hash_signature, issued_date, expiry_date', 'Stores issued-certificate details used to create the PDF and verification response.'],
    ['SMS Notifications', 'notification_id, user_id, message, status, timestamp', 'Stores status notifications sent to residents.'],
], [1.2, 2.55, 3.3])

add_paragraph('Data and Processing Modeling', 'heading')
add_paragraph('Context Diagram', 'heading')
add_paragraph('The context diagram presents CertiGuard as one system interacting with three external entities: the Resident, the Barangay Secretary or authorized staff, and the Verifier. The resident sends registration details, a dummy ID image, and certificate requests to the system. The Secretary receives the ID review queue, resident records, and certificate applications and returns approval or rejection decisions. The verifier submits a control number or QR reference and receives only the certificate validity result and non-sensitive public information.')
add_paragraph('Figure 4: Context Diagram', 'caption')

add_paragraph('Data Flow Diagram', 'heading')
add_paragraph('The Data Flow Diagram shows that registration details and ID-upload metadata enter the user and protected ID-verification records. Once the Secretary approves the registration, the resident can create a certificate request. Approved requests are sent to the certificate-generation process, which creates the certificate record, SHA-256 signature, QR reference, and notification. The public verification process reads the certificate record and returns a validity response without exposing the resident ID or full private profile.')
add_paragraph('Figure 5: Data Flow Diagram', 'caption')

add_paragraph('System Flowchart', 'heading')
add_table(['Step', 'System Action'], [
    ['1. Registration', 'Resident enters profile information, agrees to the privacy notice, and uploads a dummy ID image.'],
    ['2. ID Storage', 'System validates image type and size, then stores it in a protected identity-file location.'],
    ['3. Secretary Review', 'Secretary opens the restricted review panel and approves or rejects the residency registration.'],
    ['4. Certificate Request', 'Approved resident selects a certificate type, states the purpose, and submits the request.'],
    ['5. Certificate Processing', 'Authorized staff approves or rejects the request and records the decision.'],
    ['6. Certificate Issuance', 'System creates the certificate PDF, control number, SHA-256 signature, QR code, and notification.'],
    ['7. Public Verification', 'Verifier scans the QR code or enters a control reference to check the certificate status.'],
], [1.55, 5.5])
add_paragraph('Figure 6: System Flowchart', 'caption')

add_paragraph('Disaster Recovery Plan', 'heading')
add_paragraph('The disaster recovery plan supports continuity of the certificate service, protection of resident information, and timely restoration after a technical or security incident. The system uses controlled access, routine backups, and a defined response process for service interruption, accidental data loss, or suspected unauthorized access.')
add_table(['Goal', 'Potential Threat', 'Response and Recovery'], [
    ['Protect private resident data and dummy ID files', 'Unauthorized access, accidental disclosure, or weak access permissions', 'Use role-based access and private storage policies; revoke access, review logs, reset credentials, and notify the responsible barangay officer.'],
    ['Maintain certificate-record integrity', 'Data-entry error, record deletion, or hash mismatch', 'Restore the latest verified backup, audit the affected record, regenerate only authorized certificates, and record the incident.'],
    ['Restore availability of the service', 'Internet outage, hosting interruption, or device failure', 'Use a documented manual fallback procedure, inform applicants of delays, and restore service from the latest backup when connectivity returns.'],
    ['Preserve trustworthy verification', 'Compromised QR reference or suspected forged certificate', 'Mark the certificate as invalid when necessary, investigate the control number and hash, and issue a replacement only after staff approval.'],
], [1.35, 2.25, 3.45])
add_paragraph('Table 1: Disaster Recovery Plan', 'caption')

doc.save(OUTPUT)
print(OUTPUT)
